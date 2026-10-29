#!/usr/bin/env python3
"""Convert Markdown to neutral-styled Word document (.docx).

Usage: python3 md_to_docx.py input.md output.docx
"""
import sys
import re
import yaml
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

# Neutral document colors
ACCENT_COLOR = RGBColor(0x33, 0x41, 0x55)
STYLE_DARK = RGBColor(0x33, 0x33, 0x33)
STYLE_GRAY = RGBColor(0x66, 0x66, 0x66)
STYLE_BORDER = RGBColor(0xCC, 0xCC, 0xCC)
STYLE_TABLE_ALT = RGBColor(0xF5, 0xF5, 0xF5)
STYLE_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT_NAME = "Arial"
FONT_CODE = "Consolas"


def parse_front_matter(text: str) -> tuple[dict, str]:
    """Extract YAML front matter and body from markdown."""
    if text.startswith("---"):
        parts = text.split("---", 2)
        if len(parts) >= 3:
            meta = yaml.safe_load(parts[1]) or {}
            return meta, parts[2].strip()
    return {}, text


def set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    shading.append(shd)


def style_run(run, font_name=FONT_NAME, size=Pt(11), color=STYLE_DARK, bold=False):
    """Apply consistent styling to a run."""
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.bold = bold


def add_header_footer(doc, meta: dict):
    """Add neutral-styled header and footer."""
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    # Header
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run(meta.get("title", "Document"))
    style_run(run, size=Pt(9), color=STYLE_GRAY)
    # Accent line under header
    hp.paragraph_format.space_after = Pt(6)
    pBdr = hp._element.get_or_add_pPr().makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "6",
        qn("w:space"): "1", qn("w:color"): "334155",
    })
    pBdr.append(bottom)
    hp._element.get_or_add_pPr().append(pBdr)

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    # Accent line above footer
    pBdr2 = fp._element.get_or_add_pPr().makeelement(qn("w:pBdr"), {})
    top = pBdr2.makeelement(qn("w:top"), {
        qn("w:val"): "single", qn("w:sz"): "6",
        qn("w:space"): "1", qn("w:color"): "334155",
    })
    pBdr2.append(top)
    fp._element.get_or_add_pPr().append(pBdr2)
    classification = meta.get("classification", "Unclassified")
    run = fp.add_run(classification)
    style_run(run, size=Pt(9), color=STYLE_GRAY)
    # Page number on right via tab
    run2 = fp.add_run("\t\t")
    style_run(run2, size=Pt(9), color=STYLE_GRAY)
    # Add page number field
    fldChar1 = run2._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run2._element.append(fldChar1)
    run3 = fp.add_run()
    instrText = run3._element.makeelement(qn("w:instrText"), {})
    instrText.text = " PAGE "
    run3._element.append(instrText)
    run4 = fp.add_run()
    fldChar2 = run4._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run4._element.append(fldChar2)


def add_table(doc, header_row: list[str], data_rows: list[list[str]]):
    """Add an neutral-styled table."""
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(data_rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for i, text in enumerate(header_row):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "334155")
        p = cell.paragraphs[0]
        run = p.add_run(text.strip())
        style_run(run, size=Pt(10), color=STYLE_WHITE, bold=True)

    # Data rows
    for r, row_data in enumerate(data_rows):
        for c, text in enumerate(row_data):
            if c >= cols:
                break
            cell = table.rows[r + 1].cells[c]
            if r % 2 == 1:
                set_cell_shading(cell, "F5F5F5")
            p = cell.paragraphs[0]
            run = p.add_run(text.strip())
            style_run(run, size=Pt(10))


def parse_table_block(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    """Parse markdown table lines into header and data rows."""
    if len(lines) < 2:
        return [], []
    header = [c.strip() for c in lines[0].strip("|").split("|")]
    data = []
    for line in lines[2:]:  # skip separator
        row = [c.strip() for c in line.strip("|").split("|")]
        data.append(row)
    return header, data


def convert(input_path: str, output_path: str):
    """Convert markdown file to neutral-styled docx."""
    text = Path(input_path).read_text(encoding="utf-8")
    meta, body = parse_front_matter(text)

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(11)
    style.font.color.rgb = STYLE_DARK
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    # Heading styles
    for level, (size, color) in enumerate(
        [(Pt(24), ACCENT_COLOR), (Pt(18), ACCENT_COLOR), (Pt(14), STYLE_DARK), (Pt(12), STYLE_DARK)],
        start=1,
    ):
        hs = doc.styles[f"Heading {level}"]
        hs.font.name = FONT_NAME
        hs.font.size = size
        hs.font.color.rgb = color
        hs.font.bold = True
        hs.paragraph_format.space_before = Pt(12)
        hs.paragraph_format.space_after = Pt(6)

    add_header_footer(doc, meta)

    # Title page
    doc.add_paragraph()  # spacer
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(meta.get("title", "Document"))
    style_run(run, size=Pt(28), color=ACCENT_COLOR, bold=True)

    if meta.get("subtitle"):
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub_p.add_run(meta["subtitle"])
        style_run(run, size=Pt(16), color=STYLE_DARK)

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for key in ["date", "author", "version", "classification"]:
        if meta.get(key):
            run = info_p.add_run(f"{key.title()}: {meta[key]}    ")
            style_run(run, size=Pt(10), color=STYLE_GRAY)

    doc.add_page_break()

    # Parse body
    lines = body.split("\n")
    i = 0
    in_code = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        # Code blocks
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line.strip().lstrip("`").strip()
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                style_run(run, font_name=FONT_CODE, size=Pt(9), color=STYLE_DARK)
                # Light gray background via shading
                shd = p._element.get_or_add_pPr().makeelement(qn("w:shd"), {
                    qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "F5F5F5",
                })
                p._element.get_or_add_pPr().append(shd)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Tables
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|[\s\-:|]+\|\s*$", lines[i + 1]):
            table_lines = []
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1
            header, data = parse_table_block(table_lines)
            if header:
                add_table(doc, header, data)
            continue

        # Headings
        m = re.match(r"^(#{1,4})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        # Checkbox lists
        if re.match(r"^- \[[ x]\] ", line):
            checked = line[3] == "x"
            text = line[6:]
            p = doc.add_paragraph(style="List Bullet")
            prefix = "☑ " if checked else "☐ "
            run = p.add_run(prefix + text)
            style_run(run, size=Pt(11))
            i += 1
            continue

        # Bullet lists
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            text = re.sub(r"^[\s]*[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(text)
            style_run(run, size=Pt(11))
            i += 1
            continue

        # Numbered lists
        m = re.match(r"^\d+\.\s+(.+)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(m.group(1))
            style_run(run, size=Pt(11))
            i += 1
            continue

        # Regular paragraph
        stripped = line.strip()
        if stripped:
            p = doc.add_paragraph()
            # Handle bold/italic inline
            parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    style_run(run, bold=True)
                elif part.startswith("*") and part.endswith("*"):
                    run = p.add_run(part[1:-1])
                    style_run(run)
                    run.italic = True
                else:
                    run = p.add_run(part)
                    style_run(run)

        i += 1

    doc.save(output_path)
    print(f"Created: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(f"Usage: {sys.argv[0]} input.md output.docx")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
